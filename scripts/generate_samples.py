"""Generate synthetic sample files (fictional people only).

    python scripts/generate_samples.py

Writes:
  samples/                         spreadsheets, CSVs, placeholder photos, ZIP and PDF
  relay_hub/static/samples/        the spreadsheet/CSV files offered for download in the app
"""
from __future__ import annotations

import csv
import datetime as dt
import random
import shutil
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "samples"
STATIC = ROOT / "relay_hub" / "static" / "samples"

FIRST = ["Aarav", "Diya", "Kabir", "Isha", "Vihaan", "Ananya", "Reyansh", "Saanvi", "Arjun", "Myra", "Rohan", "Kiara",
         "Dev", "Nisha", "Yash", "Pooja", "Aditya", "Meera", "Karan", "Riya", "Nikhil", "Tara", "Harsh", "Zoya",
         "Manav", "Sneha", "Om", "Aisha", "Parth", "Lavanya"]
LAST = ["Patel", "Shah", "Mehta", "Desai", "Joshi", "Trivedi", "Parikh", "Bhatt", "Pandya", "Modi", "Rana", "Chauhan",
        "Vyas", "Dave", "Soni"]
DESIGNATIONS = ["Associate", "Senior Associate", "Analyst", "Senior Analyst", "Team Lead", "Accountant",
                "Software Engineer", "Senior Software Engineer", "HR Executive", "Recruiter"]
DEPARTMENTS = ["Finance and Accounting", "Technology", "Operations", "Human Resources", "Customer Support"]
CLIENTS = ["Northwind Traders", "Contoso Ltd", "Fabrikam Inc", "Tailspin Toys", "Wide World Importers"]
HEADERS = ["RHC ID", "Full Name", "Designation", "Department", "Official Email", "Location", "Date of Joining",
           "Date of Birth", "Employment Status", "Last Working Date", "Assignment Status", "Client Name",
           "Assignment Start", "Assignment End", "Bench Start", "Bench End", "Notes"]


def people(n=60, seed=7):
    rnd = random.Random(seed)
    used = set()
    out = []
    for i in range(1, n + 1):
        while True:
            name = f"{rnd.choice(FIRST)} {rnd.choice(LAST)}"
            if name not in used:
                used.add(name)
                break
        # A few IDs are all-digit with leading zeros to exercise text handling.
        rhc = f"00{700 + i}" if i % 15 == 0 else f"RHC{1000 + i}"
        loc = "Ahmedabad" if i % 3 else "Vadodara"
        doj = dt.date(2018, 1, 1) + dt.timedelta(days=rnd.randint(0, 2600))
        dob = dt.date(1975, 1, 1) + dt.timedelta(days=rnd.randint(0, 9000))
        status = rnd.choice(["Client", "Client", "Client", "Bench", "Internal"])
        rec = {"rhc_id": rhc, "full_name": name, "designation": rnd.choice(DESIGNATIONS),
               "department": rnd.choice(DEPARTMENTS),
               "email": name.lower().replace(" ", ".") + "@example.invalid", "location": loc, "doj": doj, "dob": dob,
               "employment": "Active", "lwd": None, "assignment": status, "client": None, "a_start": None,
               "a_end": None, "b_start": None, "b_end": None, "notes": ""}
        if status == "Client":
            rec["client"] = rnd.choice(CLIENTS)
            rec["a_start"] = doj + dt.timedelta(days=rnd.randint(10, 200))
        elif status == "Bench":
            rec["b_start"] = dt.date(2026, 9, 1) + dt.timedelta(days=rnd.randint(0, 20))
        out.append(rec)
    return out


def row(p):
    return [p["rhc_id"], p["full_name"], p["designation"], p["department"], p["email"], p["location"], p["doj"],
            p["dob"], p["employment"], p["lwd"], p["assignment"], p["client"], p["a_start"], p["a_end"], p["b_start"],
            p["b_end"], p["notes"]]


def styled_sheet(ws, headers, title=None):
    if title:
        ws.append([title])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append(["Synthetic sample data - fictional people only"])
        ws.append([])
    ws.append(headers)
    hr = ws.max_row
    for c in ws[hr]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="3B4A9C")
    return hr


def text_id_column(ws, header_row, col=1):
    for r in range(header_row + 1, ws.max_row + 1):
        cell = ws.cell(r, col)
        cell.number_format = "@"
        cell.data_type = "s"


def write_xlsx(path, rows, title=None, sheet="Employees"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    hr = styled_sheet(ws, HEADERS, title)
    for p in rows:
        ws.append(row(p))
    for r in range(hr + 1, ws.max_row + 1):
        for c in (7, 8, 10, 13, 14, 15, 16):
            ws.cell(r, c).number_format = "dd-mmm-yyyy"
    text_id_column(ws, hr)
    for i, w in enumerate([10, 22, 24, 22, 34, 12, 14, 14, 16, 16, 16, 22, 14, 14, 14, 14, 20], start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    wb.save(path)


def placeholder(path: Path, text: str, seed: int, fmt="JPEG"):
    rnd = random.Random(seed)
    bg = (rnd.randint(60, 200), rnd.randint(60, 200), rnd.randint(120, 230))
    im = Image.new("RGB", (400, 480), bg)
    d = ImageDraw.Draw(im)
    d.ellipse((120, 80, 280, 240), fill=(245, 225, 205))
    d.rounded_rectangle((80, 260, 320, 480), radius=60, fill=(40, 50, 90))
    try:
        font = ImageFont.load_default(size=36)
    except TypeError:
        font = ImageFont.load_default()
    d.text((200, 30), text, fill="white", anchor="mm", font=font)
    path.parent.mkdir(parents=True, exist_ok=True)
    im.save(path, fmt, **({"quality": 88} if fmt == "JPEG" else {}))


def make_pdf(path: Path, entries: list[tuple[str, str, Path]]):
    """A 'staff sheet' PDF: two photos per page with name and RHC ID printed under each."""
    import pypdfium2 as pdfium
    import pypdfium2.raw as raw
    import ctypes

    pdf = pdfium.PdfDocument.new()
    font = raw.FPDFText_LoadStandardFont(pdf, b"Helvetica")
    for start in range(0, len(entries), 2):
        page = pdf.new_page(595, 842)
        for slot, (name, rhc, img_path) in enumerate(entries[start:start + 2]):
            x = 70 + slot * 260
            image = pdfium.PdfImage.new(pdf)
            image.load_jpeg(str(img_path), inline=False)
            image.set_matrix(pdfium.PdfMatrix().scale(200, 240).translate(x, 480))
            page.insert_obj(image)
            for dy, text in ((0, name), (-18, f"RHC ID: {rhc}")):
                obj = raw.FPDFPageObj_NewTextObj(pdf, b"Helvetica", ctypes.c_float(12))
                buf = (text + "\x00").encode("utf-16-le")
                raw.FPDFText_SetText(obj, ctypes.cast(ctypes.c_char_p(buf), ctypes.POINTER(raw.FPDF_WCHAR)))
                raw.FPDFPageObj_Transform(obj, 1, 0, 0, 1, x, 455 + dy)
                raw.FPDFPage_InsertObject(page.raw, obj)
        page.gen_content()
    # Page with one full-page image and no text: behaves like a scanned sheet.
    page = pdf.new_page(595, 842)
    scan = pdfium.PdfImage.new(pdf)
    scan.load_jpeg(str(entries[0][2]), inline=False)
    scan.set_matrix(pdfium.PdfMatrix().scale(595, 842))
    page.insert_obj(scan)
    page.gen_content()
    _ = font
    pdf.save(str(path))


def main():
    OUT.mkdir(exist_ok=True)
    STATIC.mkdir(parents=True, exist_ok=True)
    ppl = people()

    # Blank template.
    wb = Workbook()
    ws = wb.active
    ws.title = "Employees"
    styled_sheet(ws, HEADERS)
    ws.column_dimensions["A"].number_format = "@"
    for r in range(2, 502):
        ws.cell(r, 1).number_format = "@"
    wb.save(OUT / "employee_import_template.xlsx")

    # Historical: first 45 people plus 10 former employees who left.
    hist = [dict(p) for p in ppl[:45]]
    rnd = random.Random(3)
    former = people(10, seed=99)
    for i, p in enumerate(former):
        p["rhc_id"] = f"RHC{800 + i}"
        p["email"] = p["email"].replace("@", f".old{i}@")
        p["full_name"] = p["full_name"] + " " + ["K", "M", "R", "S", "T", "V", "N", "P", "D", "B"][i]
        p["employment"] = "Left"
        p["lwd"] = dt.date(2025, 1, 1) + dt.timedelta(days=rnd.randint(0, 300))
        p["assignment"] = "Unknown"
        p["client"] = p["a_start"] = p["b_start"] = None
    for p in hist:
        p["designation"] = "Associate"  # older data: designation later changes
    write_xlsx(OUT / "sample_historical_employees.xlsx", hist + former)

    # Current roster: people 1..50 (a few new joiners appear, some historical-only people are absent).
    roster = [dict(p) for p in ppl[:50]]
    roster[4]["designation"] = "Team Lead"
    write_xlsx(OUT / "sample_current_roster.xlsx", roster, title="Relay Human Cloud - Current Roster (September 2026)",
               sheet="Roster")

    with open(OUT / "sample_partial_update.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["RHC ID", "Department", "Assignment Status", "Client Name", "Bench Start"])
        w.writerow([ppl[0]["rhc_id"], "Operations", "Bench", "", "15/09/2026"])
        w.writerow([ppl[1]["rhc_id"], "", "Client", "Contoso Ltd", ""])
        w.writerow([ppl[14]["rhc_id"], "Technology", "", "", ""])
    with open(OUT / "sample_designation_update.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["Employee ID", "Name", "New Designation"])
        for p in ppl[5:10]:
            w.writerow([p["rhc_id"], p["full_name"], "Senior " + p["designation"].replace("Senior ", "")])

    photos = OUT / "sample_photos"
    if photos.exists():
        shutil.rmtree(photos)
    for i, p in enumerate(ppl[:12]):
        if i < 6:
            fname = f"{p['rhc_id']}_{p['full_name'].replace(' ', '_')}.jpg"
        elif i < 9:
            fname = f"{p['full_name']}.jpg"
        else:
            fname = f"IMG_{4000 + i}.jpg"
        placeholder(photos / fname, p["full_name"].split()[0], i)
    placeholder(photos / "logo_banner.png", "PNG", 99, fmt="PNG")
    with open(OUT / "photo_mapping_template.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["filename", "RHC ID"])
        for i, p in enumerate(ppl[9:12]):
            w.writerow([f"IMG_{4009 + i}.jpg", p["rhc_id"]])
    with zipfile.ZipFile(OUT / "sample_photos.zip", "w") as zf:
        for f in sorted(photos.iterdir()):
            zf.write(f, f"headshots/{f.name}")
    pdf_imgs = OUT / "_pdf_imgs"
    pdf_imgs.mkdir(exist_ok=True)
    entries = []
    for i, p in enumerate(ppl[12:16]):
        img = pdf_imgs / f"{i}.jpg"
        placeholder(img, p["full_name"].split()[0], 50 + i)
        entries.append((p["full_name"], p["rhc_id"], img))
    make_pdf(OUT / "sample_headshots.pdf", entries)
    shutil.rmtree(pdf_imgs)

    for name in ("employee_import_template.xlsx", "sample_historical_employees.xlsx", "sample_current_roster.xlsx",
                 "sample_partial_update.csv", "sample_designation_update.csv", "photo_mapping_template.csv"):
        shutil.copyfile(OUT / name, STATIC / name)
    print(f"Samples written to {OUT} and {STATIC}")


if __name__ == "__main__":
    main()
