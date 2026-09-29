import datetime as dt
import io

from openpyxl import load_workbook

from relay_hub import timeutil


def test_create_read_update_archive(env):
    e = env.employee(rhc_id="rhc0001", full_name="Asha Patel", location="Ahmedabad", date_of_joining="01/02/2024",
                     official_email="Asha.Patel@Example.com")
    assert e["rhc_id"] == "RHC0001"  # normalised upper-case text
    assert e["date_of_joining"] == "2024-02-01"  # day-first
    assert e["official_email"] == "asha.patel@example.com"
    assert e["employment_status"] == "Active" and e["assignment_status"] == "Unknown"
    prof = env.api.get(f"/api/employees/{e['id']}").json()
    assert prof["employee"]["full_name"] == "Asha Patel"
    r = env.api.patch(f"/api/employees/{e['id']}", {"version": e["version"], "changes": {"designation": "Analyst"}})
    assert r.status_code == 200 and r.json()["designation"] == "Analyst" and r.json()["version"] == 2
    r = env.api.post(f"/api/employees/{e['id']}/archive", {"version": 2, "last_working_date": "2026-09-01"})
    assert r.json()["employment_status"] == "Left"
    active = env.api.get("/api/employees").json()
    assert active["total"] == 0, "archived employees leave the default (active) view"
    left = env.api.get("/api/employees?view=left").json()
    assert left["total"] == 1
    prof = env.api.get(f"/api/employees/{e['id']}").json()
    assert any(h["kind"] == "employment" for h in prof["history"])
    assert {c["field"] for c in prof["changes"]} >= {"designation", "employment_status", "last_working_date"}
    r = env.api.post(f"/api/employees/{e['id']}/reactivate", {"version": 3})
    assert r.json()["employment_status"] == "Active"


def test_leading_zeros_preserved(env):
    e = env.employee(rhc_id="000123", full_name="Zero Lead")
    assert e["rhc_id"] == "000123"
    found = env.api.get("/api/employees?search=000123").json()
    assert found["items"][0]["rhc_id"] == "000123"


def test_duplicate_rhc_and_email_rejected(env):
    env.employee(rhc_id="RHC1", full_name="A", official_email="a@x.com")
    r = env.api.post("/api/employees", {"values": {"rhc_id": "rhc1", "full_name": "B"}})
    assert r.status_code == 422 and "rhc_id" in r.json()["details"]["fields"]
    r = env.api.post("/api/employees", {"values": {"rhc_id": "RHC2", "full_name": "B", "official_email": "A@X.com"}})
    assert r.status_code == 422 and "official_email" in r.json()["details"]["fields"]


def test_version_conflict_prevents_silent_overwrite(env):
    e = env.employee(rhc_id="RHC5", full_name="Conflict Case")
    ok = env.api.patch(f"/api/employees/{e['id']}", {"version": 1, "changes": {"designation": "First"}})
    assert ok.status_code == 200
    stale = env.api.patch(f"/api/employees/{e['id']}", {"version": 1, "changes": {"designation": "Second"}})
    assert stale.status_code == 409
    assert stale.json()["details"]["current"]["designation"] == "First"
    assert env.api.get(f"/api/employees/{e['id']}").json()["employee"]["designation"] == "First"


def test_date_range_validation(env):
    r = env.api.post("/api/employees", {"values": {"rhc_id": "R1", "full_name": "X", "bench_start": "2026-05-10",
                                                   "bench_end": "2026-05-01"}})
    assert r.status_code == 422 and "bench_end" in r.json()["details"]["fields"]
    r = env.api.post("/api/employees", {"values": {"rhc_id": "R1", "full_name": "X", "location": "Surat"}})
    assert r.status_code == 422 and "location" in r.json()["details"]["fields"]
    r = env.api.post("/api/employees", {"values": {"rhc_id": "R1", "full_name": "X", "date_of_joining": "31/31/2020"}})
    assert r.status_code == 422


def test_bench_days_calculation():
    on = dt.date(2026, 9, 29)
    assert timeutil.bench_days("Bench", "2026-09-29", None, on) == 0
    assert timeutil.bench_days("Bench", "2026-09-19", None, on) == 10
    assert timeutil.bench_days("Client", "2026-09-19", None, on) is None
    assert timeutil.bench_days("Bench", "2026-10-01", None, on) is None
    assert timeutil.bench_days("Bench", "2026-09-01", "2026-09-10", on) is None
    assert timeutil.bench_days("Bench", "2026-09-01", "2026-12-31", on) == 28


def test_bench_days_in_api_and_assignment_history(env):
    timeutil.freeze(dt.datetime(2026, 9, 29, 12, 0, tzinfo=dt.timezone.utc))
    e = env.employee(rhc_id="B1", full_name="Bench Person", assignment_status="Client", client_name="Contoso",
                     assignment_start="2025-01-01")
    r = env.api.patch(f"/api/employees/{e['id']}", {"version": 1, "changes": {
        "assignment_status": "Bench", "assignment_end": "2026-09-18", "bench_start": "2026-09-19"}})
    assert r.status_code == 200, r.text
    assert r.json()["bench_days"] == 10
    hist = env.api.get(f"/api/employees/{e['id']}").json()["history"]
    assert hist[0]["kind"] == "assignment" and hist[0]["client_name"] == "Contoso"


def test_list_search_sort_filter_paginate(env):
    for i in range(12):
        env.employee(rhc_id=f"P{i:02d}", full_name=f"Person {i:02d}", location="Vadodara" if i % 2 else "Ahmedabad",
                     department="Tech" if i < 4 else "Ops")
    page = env.api.get("/api/employees?page_size=5&page=2&sort=rhc_id&direction=desc").json()
    assert page["total"] == 12 and len(page["items"]) == 5 and page["items"][0]["rhc_id"] == "P06"
    assert env.api.get("/api/employees?location=Vadodara").json()["total"] == 6
    flt = '[{"field":"department","op":"equals","value":"tech"}]'
    assert env.api.get(f"/api/employees?filters={flt}").json()["total"] == 4
    assert env.api.get("/api/employees?search=person 03").json()["total"] == 1


def test_multirow_paste_validates_then_applies_atomically(env):
    a = env.employee(rhc_id="PA", full_name="Paste A")
    b = env.employee(rhc_id="PB", full_name="Paste B")
    rows = [{"id": a["id"], "version": 1, "changes": {"designation": "Lead", "location": "Vadodara"}},
            {"id": b["id"], "version": 1, "changes": {"designation": "Dev", "location": "Nowhere"}}]
    v = env.api.post("/api/employees/paste/validate", {"rows": rows}).json()
    assert v["blocking"] == 1
    r = env.api.post("/api/employees/paste/apply", {"rows": rows})
    assert r.status_code == 422
    assert env.api.get(f"/api/employees/{a['id']}").json()["employee"]["designation"] is None, "nothing applied"
    rows[1]["changes"]["location"] = "Ahmedabad"
    r = env.api.post("/api/employees/paste/apply", {"rows": rows})
    assert r.status_code == 200 and r.json()["updated"] == 2


def test_export_current_view_protects_against_formula_injection(env):
    env.employee(rhc_id="007", full_name="=HYPERLINK(\"http://evil\")", notes="@SUM(1)", location="Ahmedabad")
    env.employee(rhc_id="008", full_name="Other City", location="Vadodara")
    r = env.api.post("/api/employees/export", {"query": {"location": "Ahmedabad"}, "columns": ["rhc_id", "full_name", "notes"]})
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("RHC ID", "Full name", "Notes")
    assert len(rows) == 2, "only the filtered view is exported"
    assert rows[1][0] == "007"
    assert rows[1][1].startswith("'=") and rows[1][2].startswith("'@")
    assert all(c.data_type == "s" for c in ws[2])
