from conftest import ADMIN, PASSWORD, upload_photo


def test_first_run_setup_only_once_and_no_default_login(env):
    c = env.new_client()
    assert c.get("/api/health").json()["setup_required"] is False
    r = c.post("/api/setup", {"email": "x@y.com", "display_name": "X", "password": "AnotherPass1!"})
    assert r.status_code == 409
    assert c.post("/api/auth/login", {"email": "admin", "password": "admin"}).status_code == 401


def test_unauthenticated_and_csrf_and_origin_checks(env):
    anon = env.new_client()
    assert anon.get("/api/employees").status_code == 401
    # Missing CSRF token on a mutating request by a signed-in user.
    r = env.api.c.post("/api/employees", json={"values": {"full_name": "X", "rhc_id": "X1"}},
                       headers={"X-Requested-With": "RelayHub"})
    assert r.status_code == 403
    # Missing application header (plain cross-site form post).
    r = env.api.c.post("/api/employees", json={"values": {}}, headers={"X-CSRF-Token": env.api.csrf})
    assert r.status_code == 403
    r = env.api.post("/api/employees", {"values": {"full_name": "X", "rhc_id": "X1"}},
                     headers={"Origin": "http://evil.example"})
    assert r.status_code == 403
    # DNS-rebinding style host header.
    r = env.api.c.get("/api/health", headers={"Host": "attacker.example"})
    assert r.status_code == 400
    r = env.api.get("/")
    assert "script-src 'self'" in r.headers["content-security-policy"]


def test_login_lockout_and_password_reset(env):
    env.add_user("ed@relay.test", "editor")
    c = env.new_client()
    for _ in range(5):
        assert c.post("/api/auth/login", {"email": "ed@relay.test", "password": "wrong-password"}).status_code == 401
    r = c.post("/api/auth/login", {"email": "ed@relay.test", "password": PASSWORD + "x"})
    assert r.status_code == 423, "locked after repeated failures even with the right password"
    uid = [u for u in env.api.get("/api/users").json() if u["email"] == "ed@relay.test"][0]["id"]
    assert env.api.post(f"/api/users/{uid}/reset-password", {"temporary_password": "TempPassw0rd!"}).status_code == 200
    c2 = env.new_client().login("ed@relay.test", "TempPassw0rd!")
    assert c2.get("/api/employees").status_code == 403, "must change the temporary password first"
    assert c2.post("/api/me/password", {"current_password": "TempPassw0rd!", "new_password": "BrandNewPass1"}).status_code == 200
    assert c2.get("/api/employees").status_code == 200


def test_role_permissions(env):
    e = env.employee(rhc_id="R1", full_name="Roled", date_of_birth="1990-03-12")
    viewer = env.add_user("view@relay.test", "viewer")
    assert viewer.get("/api/employees").status_code == 200
    assert viewer.patch(f"/api/employees/{e['id']}", {"version": 1, "changes": {"designation": "X"}}).status_code == 403
    assert viewer.post("/api/employees/export", {"query": {}}).status_code == 403
    assert viewer.post("/api/imports/upload", files={"file": ("a.csv", b"RHC ID\nX", "text/csv")}).status_code == 403
    assert viewer.get("/api/tasks").status_code == 403
    assert viewer.get("/api/users").status_code == 403
    # Export only when separately granted.
    env.add_user("view2@relay.test", "viewer", perm_grants=["employees.export"])
    v2 = env.new_client().login("view2@relay.test", PASSWORD + "x")
    assert v2.post("/api/employees/export", {"query": {}}).status_code == 200
    op = env.add_user("ops@relay.test", "export_operator")
    assert op.get("/api/tasks").status_code == 200
    assert op.post("/api/employees/export", {"query": {}}).status_code == 200
    assert op.patch(f"/api/employees/{e['id']}", {"version": 1, "changes": {"designation": "X"}}).status_code == 403
    editor = env.add_user("edit@relay.test", "editor")
    assert editor.patch(f"/api/employees/{e['id']}", {"version": 1, "changes": {"designation": "X"}}).status_code == 200
    assert editor.get("/api/settings").status_code == 403
    assert editor.get("/api/backups").status_code == 403


def test_date_of_birth_restrictions(env):
    e = env.employee(rhc_id="D1", full_name="Dob Person", date_of_birth="1990-03-12")
    viewer = env.add_user("v@relay.test", "viewer")
    item = viewer.get("/api/employees").json()["items"][0]
    assert item["date_of_birth"] == "--03-12", "viewers see day and month only by default"
    op = env.add_user("o@relay.test", "export_operator")
    item = op.get("/api/employees").json()["items"][0]
    assert "date_of_birth" not in item
    r = op.post("/api/employees/export", {"query": {}, "columns": ["rhc_id", "date_of_birth"]})
    from openpyxl import load_workbook
    import io
    header = next(load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True))
    assert "Date of birth" not in header
    editor = env.add_user("e@relay.test", "editor", dob_access="day_month")
    r = editor.patch(f"/api/employees/{e['id']}", {"version": 1, "changes": {"date_of_birth": "1991-01-01"}})
    assert r.status_code == 422, "partial DOB access cannot edit the full date"
    hidden = env.add_user("h@relay.test", "editor", hidden_fields=["notes", "official_email"])
    item = hidden.get("/api/employees").json()["items"][0]
    assert "notes" not in item and "official_email" not in item


def test_location_scope_enforced_everywhere(env):
    amd = env.employee(rhc_id="A1", full_name="In Ahmedabad", location="Ahmedabad")
    vad = env.employee(rhc_id="V1", full_name="In Vadodara", location="Vadodara")
    pid = upload_photo(env, vad["id"])
    ed = env.add_user("amd@relay.test", "editor", locations=["Ahmedabad"])
    items = ed.get("/api/employees").json()["items"]
    assert [i["rhc_id"] for i in items] == ["A1"]
    assert ed.get(f"/api/employees/{vad['id']}").status_code == 404
    assert ed.patch(f"/api/employees/{vad['id']}", {"version": 1, "changes": {"notes": "x"}}).status_code == 404
    assert ed.get(f"/api/photos/{pid}/file").status_code == 404
    assert ed.get(f"/api/photos/{pid}/thumb").status_code == 404
    r = ed.patch(f"/api/employees/{amd['id']}", {"version": 1, "changes": {"location": "Vadodara"}})
    assert r.status_code == 422, "cannot move someone out of scope"
    r = ed.post("/api/employees", {"values": {"rhc_id": "N1", "full_name": "New", "location": "Vadodara"}})
    assert r.status_code == 422


def test_fetch_tokens_scoped_hashed_revocable(env):
    env.employee(rhc_id="T1", full_name="Token A", location="Ahmedabad", date_of_birth="1990-01-01",
                 designation="Dev", notes="private")
    env.employee(rhc_id="T2", full_name="Token V", location="Vadodara")
    r = env.api.post("/api/tokens", {"name": "Display", "fields": ["designation", "date_of_birth"],
                                     "locations": ["Ahmedabad"]})
    token = r.json()["token"]
    assert "date_of_birth" not in r.json()["fields"], "restricted fields are never granted"
    with env.conn() as conn:
        stored = conn.execute("SELECT token_hash FROM api_tokens").fetchone()[0]
    assert token not in stored
    c = env.new_client()
    data = c.get("/api/fetch/v1/employees", headers={"Authorization": f"Bearer {token}"}).json()
    assert data["count"] == 1
    rec = data["employees"][0]
    assert rec["designation"] == "Dev" and "date_of_birth" not in rec and "notes" not in rec
    # A token cannot use the interactive API.
    assert c.get("/api/employees", headers={"Authorization": f"Bearer {token}"}).status_code == 401
    tid = env.api.get("/api/tokens").json()[0]["id"]
    env.api.post(f"/api/tokens/{tid}/revoke")
    assert c.get("/api/fetch/v1/employees", headers={"Authorization": f"Bearer {token}"}).status_code == 401


def test_custom_fields_everywhere(env):
    for label, ftype, opts in (("Blood group", "dropdown", ["A+", "B+", "O+"]), ("Laptop issued", "yesno", []),
                               ("Shift allowance", "number", []), ("Visa expiry", "date", []),
                               ("Emergency contact", "text", [])):
        r = env.api.post("/api/custom-fields", {"label": label, "field_type": ftype, "options": opts,
                                                "is_restricted": label == "Emergency contact"})
        assert r.status_code == 200, r.text
    keys = {f["key"] for f in env.api.get("/api/me").json()["fields"]}
    assert {"cf_blood_group", "cf_laptop_issued", "cf_shift_allowance", "cf_visa_expiry"} <= keys
    e = env.employee(rhc_id="C1", full_name="Custom", cf_blood_group="o+", cf_laptop_issued="y",
                     cf_shift_allowance="1,500", cf_visa_expiry="15/01/2027", cf_emergency_contact="Mom")
    assert e["cf_blood_group"] == "O+" and e["cf_laptop_issued"] == "Yes" and e["cf_shift_allowance"] == 1500
    assert e["cf_visa_expiry"] == "2027-01-15"
    bad = env.api.post("/api/employees", {"values": {"rhc_id": "C2", "full_name": "Bad", "cf_blood_group": "Z"}})
    assert bad.status_code == 422
    targets = {t["key"] for t in env.api.get("/api/imports/meta").json()["targets"]}
    assert "cf_blood_group" in targets
    r = env.api.post("/api/employees/export", {"query": {}, "columns": ["rhc_id", "cf_blood_group"]})
    from openpyxl import load_workbook
    import io
    rows = list(load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True))
    assert rows[0] == ("RHC ID", "Blood group") and rows[1] == ("C1", "O+")
    tok = env.api.post("/api/tokens", {"name": "t", "fields": ["cf_blood_group", "cf_emergency_contact"]}).json()
    assert "cf_emergency_contact" not in tok["fields"]
    flt = '[{"field":"cf_shift_allowance","op":"gt","value":"1000"}]'
    assert env.api.get(f"/api/employees?filters={flt}").json()["total"] == 1


def test_audit_log_records_actions(env):
    env.employee(rhc_id="AU1", full_name="Audited")
    env.add_user("aud@relay.test", "viewer")
    actions = {a["action"] for a in env.api.get("/api/audit?limit=200").json()["items"]}
    assert {"user.setup_admin", "employee.create", "user.create", "auth.login"} <= actions
    assert ADMIN[0]
